"""Reapprovisionnement trace de la buvette (09/10/2026).

Contrat : la quantite apportee s'AJOUTE au stock (15 + 60 = 75), prix d'achat
unitaire obligatoire depuis l'ecran web, dernier prix memorise sur le produit ;
la tablette garde son contrat ({product_id, delta}) et trace un reappro
« tablette » sans prix ; le PATCH d'un produit refuse `quantity` ; historique
filtre + export ; l'export d'inventaire montre les mouvements de la periode et un
recapitulatif par produit (precedent + reappros - ventes = attendu).
"""

from __future__ import annotations

import io
import uuid
from datetime import datetime, timedelta

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.crud import buvette as buvette_crud
from app.db.models import BuvetteProduct, BuvetteReapprovisionnement


pytestmark = pytest.mark.integration

CLE = "cle-de-test-de-la-caisse-au-moins-32-caracteres"
API = "/api/v1/buvette"
MESSAGE = "Le stock se change par un réapprovisionnement ou un inventaire."


@pytest.fixture(autouse=True)
def _caisse_activee(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(settings, "caisse_api_key", CLE)


@pytest.fixture()
def admin(client_authenticated_as, admin_benevoles_user) -> TestClient:
    return client_authenticated_as(admin_benevoles_user)


def _creer(admin: TestClient, nom: str, quantite: int, prix: int = 150) -> dict:
    reponse = admin.post(
        f"{API}/products",
        json={
            "name": nom,
            "price_cents": prix,
            "quantity": quantite,
            "seuil_alerte": 2,
            "caisse_category": "sucre_sale",
        },
    )
    assert reponse.status_code == 201, reponse.text
    return reponse.json()


def _reappro(admin: TestClient, produit_id: int, **corps) -> dict:
    reponse = admin.post(f"{API}/products/{produit_id}/reappro", json=corps)
    assert reponse.status_code == 200, reponse.text
    return reponse.json()


def _quantite(db: Session, produit_id: int) -> int:
    db.expire_all()
    return db.get(BuvetteProduct, produit_id).quantity


def _vente(client: TestClient, lignes: list[tuple[dict, int]]) -> None:
    corps = {
        "transaction_id": str(uuid.uuid4()),
        "sumup_tx_code": None,
        "total_cents": sum(q * p["price_cents"] for p, q in lignes),
        "sold_at": datetime.now(buvette_crud.PARIS).isoformat(),
        "lines": [
            {"product_id": p["id"], "name": p["name"], "quantity": q, "unit_price_cents": p["price_cents"]}
            for p, q in lignes
        ],
    }
    reponse = client.post(f"{API}/caisse/ventes", json=corps, headers={"X-Caisse-Key": CLE})
    assert reponse.status_code == 201, reponse.text


def _classeur(reponse):
    assert reponse.status_code == 200, reponse.text
    return load_workbook(io.BytesIO(reponse.content))


def _valeurs(feuille) -> list[tuple]:
    return [tuple(ligne) for ligne in feuille.iter_rows(values_only=True)]


# ---------------------------------------------------------------------------
# Reappro depuis l'ecran web
# ---------------------------------------------------------------------------


def test_le_reappro_s_ajoute_au_stock_et_se_trace(admin, admin_benevoles_user) -> None:
    produit = _creer(admin, "Canette", 15)

    resultat = _reappro(
        admin, produit["id"], quantite=60, prix_achat_unitaire_cents=85, commentaire="  Metro  "
    )

    assert resultat["produit"]["quantity"] == 75
    assert resultat["produit"]["dernier_prix_achat_cents"] == 85
    r = resultat["reappro"]
    assert (r["stock_avant"], r["stock_apres"]) == (15, 75)
    assert r["quantite"] == 60
    assert r["prix_achat_unitaire_cents"] == 85
    assert r["total_cents"] == 5100
    assert r["origine"] == "app"
    assert r["commentaire"] == "Metro"
    assert r["fait_par"] == admin_benevoles_user.full_name
    assert r["product_id"] == produit["id"] and r["nom"] == "Canette"
    assert r["created_at"].endswith("Z") or "+00:00" in r["created_at"]

    # Un second reappro repart du stock reel, pas d'une valeur lue a l'ecran.
    second = _reappro(admin, produit["id"], quantite=5, prix_achat_unitaire_cents=90)
    assert (second["reappro"]["stock_avant"], second["reappro"]["stock_apres"]) == (75, 80)
    assert second["produit"]["dernier_prix_achat_cents"] == 90
    assert second["reappro"]["commentaire"] is None


def test_le_prix_d_achat_est_obligatoire(admin, db_session: Session) -> None:
    produit = _creer(admin, "Chips", 4)
    url = f"{API}/products/{produit['id']}/reappro"

    for corps in (
        {"quantite": 10},
        {"quantite": 10, "prix_achat_unitaire_cents": None},
        {"quantite": 10, "prix_achat_unitaire_cents": -1},
        {"quantite": 10, "prix_achat_unitaire_cents": 50, "commentaire": "x" * 256},
        {"delta": 10},
    ):
        assert admin.post(url, json=corps).status_code == 422, corps
    assert _quantite(db_session, produit["id"]) == 4

    # Un prix nul est un prix : don, produit offert.
    assert _reappro(admin, produit["id"], quantite=10, prix_achat_unitaire_cents=0)["reappro"]["total_cents"] == 0


def test_le_dernier_prix_est_memorise_et_la_tablette_ne_l_efface_pas(
    admin, client: TestClient
) -> None:
    produit = _creer(admin, "Eau", 3)
    assert produit["dernier_prix_achat_cents"] is None
    _reappro(admin, produit["id"], quantite=12, prix_achat_unitaire_cents=40)

    reponse = client.post(
        f"{API}/caisse/reappro",
        json={"product_id": produit["id"], "delta": 10},
        headers={"X-Caisse-Key": CLE},
    )
    assert reponse.status_code == 200, reponse.text
    # Contrat de la tablette inchange.
    assert reponse.json() == {"id": produit["id"], "name": "Eau", "quantity": 25}

    fiche = next(p for p in admin.get(f"{API}/products").json() if p["id"] == produit["id"])
    assert fiche["dernier_prix_achat_cents"] == 40
    assert fiche["quantity"] == 25


def test_la_tablette_trace_son_reappro_sans_prix(admin, client: TestClient) -> None:
    produit = _creer(admin, "Gateau", 2)
    client.post(
        f"{API}/caisse/reappro",
        json={"product_id": produit["id"], "delta": 5},
        headers={"X-Caisse-Key": CLE},
    )

    liste = admin.get(f"{API}/reapprovisionnements", params={"origine": "tablette"}).json()
    assert liste["totaux"] == {"nb": 1, "quantite": 5, "montant_cents": 0}
    r = liste["reappros"][0]
    assert r["origine"] == "tablette"
    assert r["prix_achat_unitaire_cents"] is None and r["total_cents"] is None
    assert r["fait_par"] == "tablette"
    assert (r["stock_avant"], r["stock_apres"]) == (2, 7)


# ---------------------------------------------------------------------------
# Fiche produit : le stock ne s'y change plus
# ---------------------------------------------------------------------------


def test_le_patch_refuse_la_quantite(admin, db_session: Session) -> None:
    produit = _creer(admin, "The", 6)
    for corps in ({"quantity": 50}, {"quantity": None}, {"name": "The vert", "quantity": 7}):
        reponse = admin.patch(f"{API}/products/{produit['id']}", json=corps)
        assert reponse.status_code == 422, corps
        assert reponse.json()["code"] == "VAL_5001"
        assert reponse.json()["message"] == MESSAGE
    assert _quantite(db_session, produit["id"]) == 6

    # Les autres champs restent modifiables.
    reponse = admin.patch(f"{API}/products/{produit['id']}", json={"name": "The vert", "price_cents": 200})
    assert reponse.status_code == 200, reponse.text
    assert reponse.json()["name"] == "The vert"


def test_la_creation_trace_le_stock_initial(admin, admin_benevoles_user, db_session: Session) -> None:
    avec = _creer(admin, "Dattes", 12)
    _creer(admin, "Sans stock", 0)

    lignes = list(db_session.execute(select(BuvetteReapprovisionnement)).scalars())
    assert len(lignes) == 1
    ligne = lignes[0]
    assert ligne.buvette_product_id == avec["id"]
    assert (ligne.quantite, ligne.stock_avant, ligne.stock_apres) == (12, 0, 12)
    assert ligne.prix_achat_unitaire_cents is None and ligne.total_cents is None
    assert ligne.origine == "app"
    assert ligne.commentaire == "Stock initial"
    assert ligne.fait_par == admin_benevoles_user.full_name


def test_la_synchro_helloasso_et_l_import_ne_passent_pas_par_le_patch(db_session: Session) -> None:
    """La synchro initialise encore un stock nul : elle ne passe pas par le PATCH."""
    produit = BuvetteProduct(helloasso_tier_id=4242, name="Soda", price_cents=100, quantity=0)
    db_session.add(produit)
    db_session.commit()
    buvette_crud.sync_from_helloasso(
        db_session, [{"id": 4242, "label": "Soda", "price": 100, "currentQuantityAvailable": 9}]
    )
    assert _quantite(db_session, produit.id) == 9


# ---------------------------------------------------------------------------
# Historique et export
# ---------------------------------------------------------------------------


def test_l_historique_filtre_et_totalise(
    admin, db_session: Session, client_authenticated_as, compta_user, benevole_user
) -> None:
    a = _creer(admin, "Biscuits", 0)
    b = _creer(admin, "Jus", 0)
    _reappro(admin, a["id"], quantite=10, prix_achat_unitaire_cents=50)
    _reappro(admin, b["id"], quantite=4, prix_achat_unitaire_cents=120)
    ancien = _reappro(admin, a["id"], quantite=3, prix_achat_unitaire_cents=50)["reappro"]

    # Un reappro d'il y a 40 jours sort de la periode par defaut (30 jours).
    ligne = db_session.get(BuvetteReapprovisionnement, ancien["id"])
    ligne.created_at = datetime.utcnow() - timedelta(days=40)
    db_session.commit()

    corps = admin.get(f"{API}/reapprovisionnements").json()
    assert corps["totaux"] == {"nb": 2, "quantite": 14, "montant_cents": 980}
    # Plus recent d'abord.
    assert [r["nom"] for r in corps["reappros"]] == ["Jus", "Biscuits"]

    filtre = admin.get(f"{API}/reapprovisionnements", params={"product_id": a["id"]}).json()
    assert [r["quantite"] for r in filtre["reappros"]] == [10]

    jour_ancien = (datetime.now(buvette_crud.PARIS) - timedelta(days=40)).date()
    large = admin.get(
        f"{API}/reapprovisionnements",
        params={"debut": jour_ancien.isoformat(), "product_id": a["id"]},
    ).json()
    assert large["totaux"] == {"nb": 2, "quantite": 13, "montant_cents": 650}

    assert admin.get(f"{API}/reapprovisionnements", params={"origine": "tablette"}).json()["totaux"]["nb"] == 0
    assert admin.get(f"{API}/reapprovisionnements", params={"origine": "autre"}).status_code == 422
    assert admin.get(
        f"{API}/reapprovisionnements", params={"debut": "2026-10-09", "fin": "2026-10-01"}
    ).status_code == 422

    assert client_authenticated_as(compta_user).get(f"{API}/reapprovisionnements").status_code == 200
    assert client_authenticated_as(benevole_user).get(f"{API}/reapprovisionnements").status_code == 403


def test_l_export_des_reappros(admin) -> None:
    produit = _creer(admin, "Cafe moulu", 0)
    _reappro(admin, produit["id"], quantite=6, prix_achat_unitaire_cents=350, commentaire="Carrefour")

    jour = buvette_crud.aujourd_hui()
    reponse = admin.get(
        f"{API}/reapprovisionnements/export.xlsx",
        params={"debut": jour.isoformat(), "fin": jour.isoformat()},
    )
    classeur = _classeur(reponse)
    assert reponse.headers["content-disposition"] == (
        f'attachment; filename="reapprovisionnements-{jour.isoformat()}_{jour.isoformat()}.xlsx"'
    )
    assert classeur.sheetnames == ["Réapprovisionnements"]
    lignes = _valeurs(classeur["Réapprovisionnements"])
    assert lignes[0][:8] == (
        "Date", "Heure", "Produit", "Quantité", "Prix unitaire", "Total", "Origine", "Par",
    )
    assert lignes[1][2:7] == ("Cafe moulu", 6, 3.5, 21.0, "Application")
    assert lignes[1][-1] == "Carrefour"
    assert lignes[-1][0] == "Total (1 réapprovisionnements)" and lignes[-1][5] == 21.0
    assert not any("—" in str(c) for l in lignes for c in l if c)

    defaut = admin.get(f"{API}/reapprovisionnements/export.xlsx")
    debut = jour - timedelta(days=29)
    assert defaut.headers["content-disposition"] == (
        f'attachment; filename="reapprovisionnements-{debut.isoformat()}_{jour.isoformat()}.xlsx"'
    )


# ---------------------------------------------------------------------------
# Export d'inventaire : mouvements et recap par produit
# ---------------------------------------------------------------------------


def _inventaire(admin: TestClient, comptes: dict[str, int]) -> dict:
    inv = admin.post(f"{API}/inventaires")
    assert inv.status_code == 201, inv.text
    inv = inv.json()
    saisies = [{"id": l["id"], "quantite_comptee": comptes[l["nom"]]} for l in inv["lignes"]]
    assert admin.put(f"{API}/inventaires/{inv['id']}/lignes", json={"lignes": saisies}).status_code == 200
    valide = admin.post(f"{API}/inventaires/{inv['id']}/valider-stock")
    assert valide.status_code == 200, valide.text
    return valide.json()


def test_l_export_d_inventaire_recapitule_les_mouvements(admin, client: TestClient) -> None:
    canette = _creer(admin, "Canette", 10, prix=150)
    _vente(client, [(canette, 2)])  # avant le premier inventaire : hors periode

    premier = _inventaire(admin, {"Canette": 8})
    termine = admin.post(
        f"{API}/inventaires/{premier['id']}/terminer",
        json={"especes_comptees_cents": 300, "debut": buvette_crud.aujourd_hui().isoformat()},
    )
    assert termine.status_code == 200, termine.text

    # Periode du second inventaire : reappro +20 a 1,20 EUR, vente de 3, nouveau produit.
    _reappro(admin, canette["id"], quantite=20, prix_achat_unitaire_cents=120, commentaire="Metro")
    _vente(client, [(canette, 3)])
    nouveau = _creer(admin, "Nouveau", 5, prix=200)

    second = _inventaire(admin, {"Canette": 24, "Nouveau": 5})
    assert second["resume"]["achats_cents"] == 2400

    liste = admin.get(f"{API}/inventaires").json()
    achats = {i["id"]: i["resume"]["achats_cents"] for i in liste}
    assert achats == {premier["id"]: 0, second["id"]: 2400}

    classeur = _classeur(admin.get(f"{API}/inventaires/{second['id']}/export.xlsx"))
    assert classeur.sheetnames == [
        "Synthèse", "Écarts produits", "Ventes espèces",
        "Réapprovisionnements", "Mouvements", "Récap par produit",
    ]

    synthese = {l[0]: l[1] for l in _valeurs(classeur["Synthèse"]) if l[0]}
    assert synthese["Total des achats"] == 24.0
    assert synthese["Réapprovisionnements"] == 2  # +20 et le stock initial du nouveau
    assert isinstance(synthese["Période des mouvements : du"], datetime)

    reappros = _valeurs(classeur["Réapprovisionnements"])
    assert [l[2:4] for l in reappros[1:-1]] == [("Canette", 20), ("Nouveau", 5)]

    mouvements = _valeurs(classeur["Mouvements"])
    assert mouvements[0][:5] == ("Date", "Heure", "Type", "Produit", "Quantité")
    corps = [(l[2], l[3], l[4], l[5]) for l in mouvements[1:-1]]
    assert corps == [
        ("Réapprovisionnement", "Canette", 20, "Application"),
        ("Vente", "Canette", -3, "Espèces"),
        ("Réapprovisionnement", "Nouveau", 5, "Application"),
        ("Écart d'inventaire", "Canette", -1, None),
    ]

    recap = _valeurs(classeur["Récap par produit"])
    assert recap[0] == (
        "Produit", "Catégorie", "Prix unitaire", "Compté au précédent inventaire",
        "+ Réappros", "− Ventes", "= Stock attendu", "Compté", "Écart (unités)",
        "Valeur de l'écart",
    )
    par_nom = {l[0]: l for l in recap[1:] if l[0] in ("Canette", "Nouveau")}
    # precedent + reappros - ventes = attendu ; compte - attendu = ecart
    assert par_nom["Canette"][3:] == (8, 20, 3, 25, 24, -1, -1.5)
    assert par_nom["Nouveau"][3:] == (0, 5, 0, 5, 5, 0, 0.0)
    # Coherent avec l'ecart de l'inventaire (theorique = stock en base).
    ecart_inventaire = {l["nom"]: l["ecart"] for l in second["lignes"]}
    assert ecart_inventaire == {"Canette": -1, "Nouveau": 0}

    for nom in classeur.sheetnames:
        assert not any("—" in str(c) for l in _valeurs(classeur[nom]) for c in l if c), nom

    # Premier inventaire : aucun precedent, le recap ne devine rien.
    premier_classeur = _classeur(admin.get(f"{API}/inventaires/{premier['id']}/export.xlsx"))
    recap_premier = _valeurs(premier_classeur["Récap par produit"])
    assert recap_premier[1][0] == "Canette" and recap_premier[1][3] is None
    synthese_premier = {l[0]: l[1] for l in _valeurs(premier_classeur["Synthèse"]) if l[0]}
    assert synthese_premier["Période des mouvements : du"] == "Début des données"


def test_l_export_de_l_historique_ajoute_les_reappros(admin) -> None:
    produit = _creer(admin, "Pain", 0)
    _reappro(admin, produit["id"], quantite=8, prix_achat_unitaire_cents=100)
    jour = buvette_crud.aujourd_hui().isoformat()
    classeur = _classeur(
        admin.get(f"{API}/inventaires/export.xlsx", params={"debut": jour, "fin": jour})
    )
    assert classeur.sheetnames == ["Inventaires", "Détail", "Réapprovisionnements"]
    lignes = _valeurs(classeur["Réapprovisionnements"])
    assert lignes[1][2:6] == ("Pain", 8, 1.0, 8.0)

    # Hors periode : feuille vide (en-tetes et totaux seulement).
    vide = _classeur(
        admin.get(f"{API}/inventaires/export.xlsx", params={"debut": "2026-01-01", "fin": "2026-01-31"})
    )
    assert len(_valeurs(vide["Réapprovisionnements"])) == 2
