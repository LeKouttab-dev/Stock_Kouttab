"""Cloture de caisse especes PONCTUELLE (10/10/2026).

Contrat : on cloture quand on veut et a chaque comptage on VIDE la boite.
Attendu = ventes especes de la tablette (sans code SumUp) dans ]dernier
comptage ; maintenant], le dernier comptage etant le plus recent entre la
derniere cloture et la fin du dernier inventaire termine. Premier comptage de
tous : depuis une date choisie. L'inventaire suit la meme regle.
"""

from __future__ import annotations

import io
import uuid
from datetime import date, datetime, timedelta
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.core.config import settings
from app.crud import buvette as buvette_crud
from app.db.models import BuvetteProduct, ClotureCaisse, Inventaire
from app.schemas.buvette import BuvetteProductCreate


pytestmark = pytest.mark.integration

CLE = "cle-de-test-de-la-caisse-au-moins-32-caracteres"
API = "/api/v1/buvette"


@pytest.fixture(autouse=True)
def _caisse_activee(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(settings, "caisse_api_key", CLE)


def _produit(db: Session, nom: str, prix: int) -> BuvetteProduct:
    return buvette_crud.create_product(
        db,
        BuvetteProductCreate(
            name=nom, price_cents=prix, quantity=50, seuil_alerte=1, caisse_category="sucre_sale"
        ),
    )


def _vente(
    client: TestClient, produit: BuvetteProduct, quantite: int = 1, *, carte: bool = False,
    quand: datetime | None = None,
) -> None:
    """Vente de la tablette ; `quand` = heure de Paris (par defaut, maintenant)."""
    quand = quand or datetime.now(buvette_crud.PARIS)
    if quand.tzinfo is None:
        quand = quand.replace(tzinfo=buvette_crud.PARIS)
    corps = {
        "transaction_id": str(uuid.uuid4()),
        "sumup_tx_code": f"TX{uuid.uuid4().hex[:8]}" if carte else None,
        "total_cents": quantite * produit.price_cents,
        "sold_at": quand.isoformat(),
        "lines": [{
            "product_id": produit.id, "name": produit.name, "quantity": quantite,
            "unit_price_cents": produit.price_cents,
        }],
    }
    reponse = client.post(f"{API}/caisse/ventes", json=corps, headers={"X-Caisse-Key": CLE})
    assert reponse.status_code == 201, reponse.text


def _il_y_a(**delta: float) -> datetime:
    return datetime.now(buvette_crud.PARIS) - timedelta(**delta)


def _vieillir_clotures(db: Session, **delta: float) -> None:
    """Recule toutes les clotures dans le passe (UTC naif), pour simuler le temps."""
    db.expire_all()
    for c in db.query(ClotureCaisse).all():
        c.periode_fin = c.periode_fin - timedelta(**delta)
        c.periode_debut = c.periode_debut - timedelta(**delta)
        c.created_at = c.created_at - timedelta(**delta)
    db.commit()


@pytest.fixture()
def produits(db_session: Session) -> dict[str, BuvetteProduct]:
    return {"the": _produit(db_session, "The", 100), "gateau": _produit(db_session, "Gateau", 250)}


@pytest.fixture()
def gerant(client_authenticated_as, admin_stock_user) -> TestClient:
    return client_authenticated_as(admin_stock_user)


def HIER() -> date:  # noqa: N802
    return buvette_crud.aujourd_hui() - timedelta(days=1)


def _cloturer(api: TestClient, compte: int, **extra: Any) -> dict[str, Any]:
    reponse = api.post(f"{API}/clotures", json={"compte_cents": compte, **extra})
    assert reponse.status_code == 201, reponse.text
    return reponse.json()


# ---------------------------------------------------------------------------
# Premier comptage
# ---------------------------------------------------------------------------


def test_premier_comptage_sans_date(produits, gerant, client) -> None:
    _vente(client, produits["the"], 2, quand=_il_y_a(minutes=10))
    lecture = gerant.get(f"{API}/clotures/attendu").json()
    assert lecture["premier_comptage"] is True
    assert lecture["dernier_comptage"] is None
    assert lecture["periode_debut"] is None
    assert lecture["attendu_cents"] == 0 and lecture["ventes"] == []

    reponse = gerant.post(f"{API}/clotures", json={"compte_cents": 200})
    assert reponse.status_code == 422
    assert "Premier comptage" in reponse.text


def test_premier_comptage_avec_date(produits, gerant, client) -> None:
    _vente(client, produits["the"], 2, quand=_il_y_a(minutes=30))
    _vente(client, produits["gateau"], 1, quand=_il_y_a(minutes=20))
    _vente(client, produits["gateau"], 1, carte=True, quand=_il_y_a(minutes=15))  # carte : exclue
    _vente(client, produits["the"], 1, quand=_il_y_a(days=3))  # avant la date choisie

    lecture = gerant.get(f"{API}/clotures/attendu", params={"debut": HIER().isoformat()}).json()
    assert lecture["premier_comptage"] is True
    assert lecture["attendu_cents"] == 450 and lecture["nb_ventes"] == 2
    vente = lecture["ventes"][0]
    assert set(vente) == {"cle", "sold_at", "total_cents", "articles"}
    assert vente["sold_at"].endswith("Z") or "+00:00" in vente["sold_at"]
    assert vente["articles"][0]["nom"] == "The"

    cloture = _cloturer(gerant, 400, debut=HIER().isoformat(), commentaire="  Pièce perdue  ")
    assert cloture["attendu_cents"] == 450
    assert cloture["ecart_cents"] == -50
    assert cloture["nb_ventes"] == 2
    assert cloture["commentaire"] == "Pièce perdue"
    assert cloture["saisi_par"]
    assert cloture["jour"] == buvette_crud.aujourd_hui().isoformat()
    assert cloture["periode_debut"] and cloture["periode_fin"]


# ---------------------------------------------------------------------------
# Clotures successives
# ---------------------------------------------------------------------------


def test_deux_clotures_le_meme_jour_et_attendu_depuis_la_derniere(
    produits, gerant, client, db_session
) -> None:
    _vente(client, produits["the"], 3, quand=_il_y_a(minutes=30))
    premiere = _cloturer(gerant, 300, debut=HIER().isoformat())
    assert premiere["attendu_cents"] == 300

    _vente(client, produits["gateau"], 1)
    lecture = gerant.get(f"{API}/clotures/attendu", params={"debut": "2020-01-01"}).json()
    # La date est ignoree : la periode part de la cloture precedente.
    assert lecture["premier_comptage"] is False
    assert lecture["dernier_comptage"]["type"] == "cloture"
    assert lecture["attendu_cents"] == 250 and lecture["nb_ventes"] == 1

    seconde = _cloturer(gerant, 250)
    assert seconde["attendu_cents"] == 250 and seconde["ecart_cents"] == 0
    assert seconde["jour"] == premiere["jour"]
    assert seconde["periode_debut"] == premiere["periode_fin"]

    historique = gerant.get(f"{API}/clotures").json()
    assert [c["id"] for c in historique] == [seconde["id"], premiere["id"]]


def test_un_doublon_immediat_sans_vente_est_refuse(produits, gerant, client, db_session) -> None:
    _vente(client, produits["the"], 1, quand=_il_y_a(minutes=5))
    _cloturer(gerant, 100, debut=HIER().isoformat())

    reponse = gerant.post(f"{API}/clotures", json={"compte_cents": 100})
    assert reponse.status_code == 409
    assert "doublon" in reponse.text

    # Une minute plus tard, une cloture sans vente (boite vide) est permise.
    _vieillir_clotures(db_session, seconds=61)
    vide = _cloturer(gerant, 0)
    assert vide["attendu_cents"] == 0 and vide["nb_ventes"] == 0


def test_une_vente_entre_deux_clotures_rapprochees_n_est_pas_un_doublon(
    produits, gerant, client
) -> None:
    _cloturer(gerant, 0, debut=buvette_crud.aujourd_hui().isoformat())
    _vente(client, produits["the"], 1)
    assert _cloturer(gerant, 100)["attendu_cents"] == 100


# ---------------------------------------------------------------------------
# Avec l'inventaire
# ---------------------------------------------------------------------------


def _inventaire_termine(api: TestClient, comptees: int, **extra: Any) -> dict[str, Any]:
    inv = api.post(f"{API}/inventaires").json()
    assert api.post(f"{API}/inventaires/{inv['id']}/valider-stock").status_code == 200
    reponse = api.post(
        f"{API}/inventaires/{inv['id']}/terminer",
        json={"especes_comptees_cents": comptees, **extra},
    )
    assert reponse.status_code == 200, reponse.text
    return reponse.json()


def test_l_inventaire_apres_une_cloture_ne_recompte_pas_ses_ventes(
    produits, gerant, client
) -> None:
    _vente(client, produits["the"], 2, quand=_il_y_a(minutes=30))
    _cloturer(gerant, 200, debut=HIER().isoformat())
    _vente(client, produits["gateau"], 1)

    inv = gerant.post(f"{API}/inventaires").json()
    especes = gerant.get(f"{API}/inventaires/{inv['id']}/especes").json()
    assert especes["premier_inventaire"] is False
    assert especes["dernier_comptage"]["type"] == "cloture"
    assert especes["attendu_cents"] == 250 and especes["nb_ventes"] == 1

    gerant.post(f"{API}/inventaires/{inv['id']}/valider-stock")
    fini = gerant.post(
        f"{API}/inventaires/{inv['id']}/terminer", json={"especes_comptees_cents": 250}
    ).json()
    assert fini["especes_attendues_cents"] == 250 and fini["ecart_especes_cents"] == 0


def test_la_cloture_apres_un_inventaire_part_de_sa_fin(
    produits, gerant, client, db_session
) -> None:
    _vente(client, produits["the"], 2, quand=_il_y_a(hours=2))
    fini = _inventaire_termine(gerant, 200, debut=HIER().isoformat())
    assert fini["especes_attendues_cents"] == 200
    # L'inventaire s'est termine il y a une heure.
    db_session.expire_all()
    inventaire = db_session.get(Inventaire, fini["id"])
    inventaire.termine_le = datetime.utcnow() - timedelta(hours=1)
    db_session.commit()

    _vente(client, produits["gateau"], 1, quand=_il_y_a(minutes=10))
    lecture = gerant.get(f"{API}/clotures/attendu").json()
    assert lecture["premier_comptage"] is False
    assert lecture["dernier_comptage"]["type"] == "inventaire"
    assert lecture["attendu_cents"] == 250

    assert _cloturer(gerant, 250)["attendu_cents"] == 250


# ---------------------------------------------------------------------------
# Droits, validation, historique, export
# ---------------------------------------------------------------------------


def test_la_compta_consulte_mais_ne_cloture_pas(client_authenticated_as, compta_user) -> None:
    compta = client_authenticated_as(compta_user)
    assert compta.get(f"{API}/clotures/attendu").status_code == 200
    assert compta.get(f"{API}/clotures").status_code == 200
    reponse = compta.post(f"{API}/clotures", json={"compte_cents": 0, "debut": "2026-10-01"})
    assert reponse.status_code == 403


def test_un_compte_negatif_ou_une_date_future_sont_refuses(gerant) -> None:
    assert gerant.post(
        f"{API}/clotures", json={"compte_cents": -1, "debut": "2026-10-01"}
    ).status_code == 422
    demain = (buvette_crud.aujourd_hui() + timedelta(days=2)).isoformat()
    assert gerant.post(
        f"{API}/clotures", json={"compte_cents": 0, "debut": demain}
    ).status_code == 422


def _ancienne_cloture(db: Session, jour: date, ecart: int = 0) -> ClotureCaisse:
    debut = datetime.combine(jour, datetime.min.time())
    cloture = ClotureCaisse(
        jour=jour, periode_debut=debut, periode_fin=debut + timedelta(hours=20),
        nb_ventes=3, attendu_cents=1000, compte_cents=1000 + ecart, ecart_cents=ecart,
        commentaire=None, saisi_par="Yusuf", created_at=debut + timedelta(hours=20),
    )
    db.add(cloture)
    db.commit()
    return cloture


def test_historique_filtre_et_export(gerant, db_session) -> None:
    _ancienne_cloture(db_session, date(2026, 10, 1), ecart=-50)
    _ancienne_cloture(db_session, date(2026, 10, 3), ecart=20)
    _ancienne_cloture(db_session, date(2026, 10, 3))

    def jours(**params: str) -> list[str]:
        reponse = gerant.get(f"{API}/clotures", params=params)
        assert reponse.status_code == 200, reponse.text
        return [c["jour"] for c in reponse.json()]

    assert jours(debut="2026-10-02", fin="2026-10-05") == ["2026-10-03", "2026-10-03"]
    assert jours(debut="2026-10-01", fin="2026-10-01") == ["2026-10-01"]
    assert gerant.get(
        f"{API}/clotures", params={"debut": "2026-10-05", "fin": "2026-10-01"}
    ).status_code == 422

    reponse = gerant.get(
        f"{API}/clotures/export.xlsx", params={"debut": "2026-10-01", "fin": "2026-10-03"}
    )
    assert reponse.status_code == 200, reponse.text
    assert reponse.headers["content-disposition"] == (
        'attachment; filename="clotures-2026-10-01_2026-10-03.xlsx"'
    )
    classeur = load_workbook(io.BytesIO(reponse.content))
    assert classeur.sheetnames == ["Clôtures de caisse"]
    lignes = list(classeur["Clôtures de caisse"].iter_rows(values_only=True))
    assert lignes[0][:8] == (
        "Date", "Heure", "Période : du", "Période : au", "Ventes en espèces",
        "Attendu", "Compté", "Écart (compté moins attendu)",
    )
    assert lignes[-1][0] == "Total (3 clôtures)"
    assert lignes[-1][7] == -0.3


# ---------------------------------------------------------------------------
# Export d'inventaire : feuille « Clôtures de caisse »
# ---------------------------------------------------------------------------


def test_l_export_d_inventaire_reprend_les_clotures_de_la_periode(
    produits, gerant, client
) -> None:
    _vente(client, produits["the"], 2, quand=_il_y_a(minutes=30))
    _cloturer(gerant, 150, debut=HIER().isoformat(), commentaire="Billet manquant")
    inv = gerant.post(f"{API}/inventaires").json()
    gerant.post(f"{API}/inventaires/{inv['id']}/valider-stock")
    gerant.post(f"{API}/inventaires/{inv['id']}/terminer", json={"especes_comptees_cents": 0})

    reponse = gerant.get(f"{API}/inventaires/{inv['id']}/export.xlsx")
    assert reponse.status_code == 200, reponse.text
    classeur = load_workbook(io.BytesIO(reponse.content))
    assert "Clôtures de caisse" in classeur.sheetnames

    lignes = list(classeur["Clôtures de caisse"].iter_rows(values_only=True))
    assert len(lignes) == 1 + 1 + 1
    assert lignes[1][5:10] == (2.0, 1.5, -0.5, "Billet manquant", lignes[1][9])
    assert lignes[1][9]  # par
    assert lignes[-1][0] == "Total (1 clôtures)"

    synthese = {l[0]: l[1] for l in classeur["Synthèse"].iter_rows(values_only=True) if l[0]}
    assert synthese["Clôtures de caisse de la période"] == 1
    assert synthese["Cumul des écarts des clôtures"] == -0.5

    mouvements = list(classeur["Mouvements"].iter_rows(values_only=True))
    (cloture,) = [m for m in mouvements if m[2] == "Clôture de caisse"]
    assert cloture[5] == "Espèces" and cloture[7] == 1.5
    for nom in classeur.sheetnames:
        for ligne in classeur[nom].iter_rows(values_only=True):
            assert not any("—" in str(c) for c in ligne if c), nom
