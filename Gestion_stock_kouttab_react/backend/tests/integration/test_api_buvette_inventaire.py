"""Inventaire de la buvette (stock puis especes) et exports Excel.

Contrat : stock remplace par les quantites comptees a la validation, theorique
fige A CET INSTANT (une vente passee pendant le comptage n'est pas un manque),
cafes exclus, especes = ventes de la tablette sans code SumUp depuis la fin du
dernier inventaire termine, un seul inventaire ouvert a la fois.
"""

from __future__ import annotations

import io
import uuid
from datetime import date, datetime, timedelta
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.core.config import settings
from app.crud import buvette as buvette_crud
from app.db.models import BuvetteProduct, Inventaire
from app.schemas.buvette import BuvetteProductCreate


pytestmark = pytest.mark.integration

CLE = "cle-de-test-de-la-caisse-au-moins-32-caracteres"
API = "/api/v1/buvette"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture(autouse=True)
def _caisse_activee(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(settings, "caisse_api_key", CLE)


def _produit(
    db: Session,
    *,
    nom: str,
    quantite: int = 10,
    seuil: int = 2,
    prix: int = 150,
    categorie: str | None = "sucre_sale",
    actif: bool = True,
) -> BuvetteProduct:
    return buvette_crud.create_product(
        db,
        BuvetteProductCreate(
            name=nom,
            price_cents=prix,
            quantity=quantite,
            seuil_alerte=seuil,
            caisse_category=categorie,
            is_active=actif,
        ),
    )


def _vente(
    client: TestClient,
    lignes: list[tuple[BuvetteProduct, int]],
    *,
    carte: bool = False,
    quand: datetime | None = None,
) -> str:
    """Vente de la tablette ; `quand` = heure de Paris (par defaut, maintenant)."""
    quand = quand or datetime.now(buvette_crud.PARIS)
    if quand.tzinfo is None:
        quand = quand.replace(tzinfo=buvette_crud.PARIS)
    tx = str(uuid.uuid4())
    corps = {
        "transaction_id": tx,
        "sumup_tx_code": f"TX{uuid.uuid4().hex[:8]}" if carte else None,
        "total_cents": sum(q * p.price_cents for p, q in lignes),
        "sold_at": quand.isoformat(),
        "lines": [
            {"product_id": p.id, "name": p.name, "quantity": q, "unit_price_cents": p.price_cents}
            for p, q in lignes
        ],
    }
    reponse = client.post(f"{API}/caisse/ventes", json=corps, headers={"X-Caisse-Key": CLE})
    assert reponse.status_code == 201, reponse.text
    return tx


def _il_y_a(**delta: float) -> datetime:
    return datetime.now(buvette_crud.PARIS) - timedelta(**delta)


def _quantite(db: Session, produit_id: int) -> int:
    db.expire_all()
    return db.get(BuvetteProduct, produit_id).quantity


def _lignes_par_nom(inventaire: dict[str, Any]) -> dict[str, dict[str, Any]]:
    return {l["nom"]: l for l in inventaire["lignes"]}


def _compter(api: TestClient, inventaire: dict[str, Any], comptes: dict[str, int]) -> dict[str, Any]:
    par_nom = _lignes_par_nom(inventaire)
    reponse = api.put(
        f"{API}/inventaires/{inventaire['id']}/lignes",
        json={"lignes": [{"id": par_nom[n]["id"], "quantite_comptee": q} for n, q in comptes.items()]},
    )
    assert reponse.status_code == 200, reponse.text
    return reponse.json()


def _deplacer(db: Session, inventaire_id: int, **champs: Any) -> None:
    """Recale des horodatages (UTC naif) pour simuler le passe."""
    db.expire_all()
    inventaire = db.get(Inventaire, inventaire_id)
    for nom, valeur in champs.items():
        setattr(inventaire, nom, valeur)
    db.commit()


def _utc_naif(**delta: float) -> datetime:
    return datetime.utcnow() - timedelta(**delta)


@pytest.fixture()
def catalogue(db_session: Session) -> dict[str, BuvetteProduct]:
    return {
        "the": _produit(db_session, nom="The", quantite=10, prix=100, categorie="boissons"),
        "gateau": _produit(db_session, nom="Gateau", quantite=5, prix=250),
        "chips": _produit(db_session, nom="Chips", quantite=4, prix=120, categorie="epicerie", seuil=5),
        "cafe": _produit(db_session, nom="Cafe", quantite=50, prix=100, categorie="cafe"),
        "retire": _produit(db_session, nom="Retire", quantite=3, actif=False),
        "hors_tablette": _produit(db_session, nom="Hors tablette", quantite=3, categorie=None),
    }


@pytest.fixture()
def admin(client_authenticated_as, admin_benevoles_user) -> TestClient:
    return client_authenticated_as(admin_benevoles_user)


def _demarrer(api: TestClient) -> dict[str, Any]:
    reponse = api.post(f"{API}/inventaires")
    assert reponse.status_code == 201, reponse.text
    return reponse.json()


# ---------------------------------------------------------------------------
# Demarrage, brouillon, abandon
# ---------------------------------------------------------------------------


def test_demarrer_cree_les_lignes_a_zero_hors_cafes(catalogue, admin, admin_benevoles_user) -> None:
    inv = _demarrer(admin)
    assert inv["statut"] == "en_cours"
    assert inv["cree_par"] == admin_benevoles_user.full_name
    assert inv["debut_le"].endswith(("Z", "+00:00"))
    noms = {l["nom"] for l in inv["lignes"]}
    # Cafes, produits inactifs et produits absents de la tablette ne se comptent pas.
    assert noms == {"The", "Gateau", "Chips"}
    the = _lignes_par_nom(inv)["The"]
    assert the["quantite_comptee"] == 0
    assert the["stock_actuel"] == 10
    assert the["prix_cents"] == 100
    assert the["categorie"] == "boissons"
    assert the["product_id"] == catalogue["the"].id
    assert the["quantite_theorique"] is None and the["ecart"] is None
    assert inv["resume"] == {
        "nb_produits": 3, "nb_ecarts": 0, "ecart_unites": 0, "valeur_ecart_cents": 0, "perte_cents": 0,
    }


def test_un_seul_inventaire_ouvert_a_la_fois(catalogue, admin) -> None:
    inv = _demarrer(admin)
    assert admin.post(f"{API}/inventaires").status_code == 409
    # Toujours refuse une fois le stock valide : l'inventaire n'est pas termine.
    assert admin.post(f"{API}/inventaires/{inv['id']}/valider-stock").status_code == 200
    assert admin.post(f"{API}/inventaires").status_code == 409


def test_le_brouillon_est_repris(catalogue, admin) -> None:
    assert admin.get(f"{API}/inventaires/en-cours").json() == {"inventaire": None}
    inv = _demarrer(admin)
    _compter(admin, inv, {"The": 7, "Gateau": 5})

    repris = admin.get(f"{API}/inventaires/en-cours").json()["inventaire"]
    assert repris["id"] == inv["id"]
    par_nom = _lignes_par_nom(repris)
    assert par_nom["The"]["quantite_comptee"] == 7
    assert par_nom["Gateau"]["quantite_comptee"] == 5
    assert par_nom["Chips"]["quantite_comptee"] == 0
    assert admin.get(f"{API}/inventaires/{inv['id']}").json()["lignes"] == repris["lignes"]


def test_brouillon_refuse_une_ligne_etrangere_ou_negative(catalogue, admin) -> None:
    inv = _demarrer(admin)
    assert admin.put(
        f"{API}/inventaires/{inv['id']}/lignes", json={"lignes": [{"id": 999999, "quantite_comptee": 1}]}
    ).status_code == 422
    ligne = inv["lignes"][0]["id"]
    assert admin.put(
        f"{API}/inventaires/{inv['id']}/lignes", json={"lignes": [{"id": ligne, "quantite_comptee": -1}]}
    ).status_code == 422


def test_inventaire_introuvable(admin) -> None:
    assert admin.get(f"{API}/inventaires/424242").status_code == 404
    assert admin.post(f"{API}/inventaires/424242/valider-stock").status_code == 404


def test_abandonner_un_brouillon(catalogue, admin, db_session) -> None:
    inv = _demarrer(admin)
    _compter(admin, inv, {"The": 1})
    assert admin.delete(f"{API}/inventaires/{inv['id']}").status_code == 204
    assert admin.get(f"{API}/inventaires/{inv['id']}").status_code == 404
    assert admin.get(f"{API}/inventaires/en-cours").json() == {"inventaire": None}
    # Le stock n'a pas bouge, et un nouvel inventaire peut demarrer.
    assert _quantite(db_session, catalogue["the"].id) == 10
    _demarrer(admin)


def test_on_n_abandonne_pas_un_inventaire_dont_le_stock_est_valide(catalogue, admin) -> None:
    inv = _demarrer(admin)
    admin.post(f"{API}/inventaires/{inv['id']}/valider-stock")
    assert admin.delete(f"{API}/inventaires/{inv['id']}").status_code == 409


# ---------------------------------------------------------------------------
# Validation du stock
# ---------------------------------------------------------------------------


def test_valider_le_stock_fige_les_ecarts_et_met_le_stock_a_jour(
    catalogue, admin, client, db_session, captured_emails
) -> None:
    inv = _demarrer(admin)
    # Vente pendant le comptage : 2 thes partent, le stock en base passe a 8.
    _vente(client, [(catalogue["the"], 2)], carte=True)
    inv = _compter(admin, inv, {"The": 8, "Gateau": 3, "Chips": 6})
    assert _lignes_par_nom(inv)["The"]["stock_actuel"] == 8
    courriels_avant = len(captured_emails)

    reponse = admin.post(f"{API}/inventaires/{inv['id']}/valider-stock")
    assert reponse.status_code == 200, reponse.text
    valide = reponse.json()
    assert valide["statut"] == "stock_valide"
    assert valide["stock_valide_le"] is not None
    par_nom = _lignes_par_nom(valide)

    # La vente du comptage est dans le theorique : pas de manque sur le the.
    assert (par_nom["The"]["quantite_theorique"], par_nom["The"]["ecart"]) == (8, 0)
    assert par_nom["The"]["valeur_ecart_cents"] == 0
    assert (par_nom["Gateau"]["quantite_theorique"], par_nom["Gateau"]["ecart"]) == (5, -2)
    assert par_nom["Gateau"]["valeur_ecart_cents"] == -500
    assert (par_nom["Chips"]["quantite_theorique"], par_nom["Chips"]["ecart"]) == (4, 2)
    assert par_nom["Chips"]["valeur_ecart_cents"] == 240
    assert valide["resume"] == {
        "nb_produits": 3, "nb_ecarts": 2, "ecart_unites": 0, "valeur_ecart_cents": -260, "perte_cents": 500,
    }

    # Le stock est remplace par les quantites comptees.
    assert _quantite(db_session, catalogue["the"].id) == 8
    assert _quantite(db_session, catalogue["gateau"].id) == 3
    assert _quantite(db_session, catalogue["chips"].id) == 6
    assert _quantite(db_session, catalogue["cafe"].id) == 50
    assert par_nom["Gateau"]["stock_actuel"] == 3

    # alert_sent recale sur le seuil, sans aucun courriel.
    db_session.expire_all()
    assert db_session.get(BuvetteProduct, catalogue["chips"].id).alert_sent is False
    assert db_session.get(BuvetteProduct, catalogue["gateau"].id).alert_sent is False
    assert len(captured_emails) == courriels_avant

    # Fige : plus de brouillon, plus de seconde validation.
    ligne = par_nom["The"]["id"]
    assert admin.put(
        f"{API}/inventaires/{inv['id']}/lignes", json={"lignes": [{"id": ligne, "quantite_comptee": 1}]}
    ).status_code == 409
    assert admin.post(f"{API}/inventaires/{inv['id']}/valider-stock").status_code == 409


def test_valider_leve_alert_sent_sous_le_seuil_sans_prevenir(
    catalogue, admin, db_session, captured_emails
) -> None:
    inv = _demarrer(admin)
    _compter(admin, inv, {"The": 10, "Gateau": 5, "Chips": 1})
    admin.post(f"{API}/inventaires/{inv['id']}/valider-stock")
    db_session.expire_all()
    chips = db_session.get(BuvetteProduct, catalogue["chips"].id)
    assert chips.quantity == 1 and chips.alert_sent is True
    assert captured_emails == []


def test_un_produit_supprime_pendant_le_comptage(catalogue, admin, db_session) -> None:
    inv = _demarrer(admin)
    buvette_crud.delete_product(db_session, catalogue["gateau"].id)
    valide = admin.post(f"{API}/inventaires/{inv['id']}/valider-stock").json()
    gateau = _lignes_par_nom(valide)["Gateau"]
    assert gateau["product_id"] is None
    assert gateau["stock_actuel"] is None
    assert gateau["quantite_theorique"] is None and gateau["ecart"] is None


# ---------------------------------------------------------------------------
# Especes
# ---------------------------------------------------------------------------


def test_premier_inventaire_sans_date(catalogue, admin) -> None:
    inv = _demarrer(admin)
    especes = admin.get(f"{API}/inventaires/{inv['id']}/especes").json()
    assert especes["premier_inventaire"] is True
    assert especes["periode_debut"] is None
    assert especes["ventes"] == [] and especes["attendu_cents"] == 0 and especes["nb_ventes"] == 0
    assert especes["periode_fin"] is not None


def test_premier_inventaire_avec_date(catalogue, admin, client) -> None:
    _vente(client, [(catalogue["the"], 2), (catalogue["gateau"], 1)], quand=_il_y_a(minutes=30))
    _vente(client, [(catalogue["chips"], 1)], quand=_il_y_a(minutes=20))
    _vente(client, [(catalogue["gateau"], 1)], carte=True, quand=_il_y_a(minutes=10))
    # Avant la periode choisie : exclue.
    _vente(client, [(catalogue["the"], 1)], quand=_il_y_a(days=3))
    inv = _demarrer(admin)

    hier = buvette_crud.aujourd_hui() - timedelta(days=1)
    especes = admin.get(f"{API}/inventaires/{inv['id']}/especes", params={"debut": hier.isoformat()}).json()
    assert especes["premier_inventaire"] is True
    assert especes["periode_debut"] is not None
    assert especes["nb_ventes"] == 2
    assert especes["attendu_cents"] == 2 * 100 + 250 + 120
    # Ordre chronologique, articles detailles, heures en UTC avec fuseau.
    premiere = especes["ventes"][0]
    assert premiere["total_cents"] == 450
    assert [(a["nom"], a["quantite"], a["montant_cents"]) for a in premiere["articles"]] == [
        ("The", 2, 200), ("Gateau", 1, 250),
    ]
    assert premiere["sold_at"].endswith(("Z", "+00:00"))


def test_terminer_le_premier_inventaire_exige_une_date(catalogue, admin) -> None:
    inv = _demarrer(admin)
    admin.post(f"{API}/inventaires/{inv['id']}/valider-stock")
    reponse = admin.post(f"{API}/inventaires/{inv['id']}/terminer", json={"especes_comptees_cents": 0})
    assert reponse.status_code == 422


def test_terminer_avant_de_valider_le_stock(catalogue, admin) -> None:
    inv = _demarrer(admin)
    reponse = admin.post(
        f"{API}/inventaires/{inv['id']}/terminer",
        json={"especes_comptees_cents": 0, "debut": buvette_crud.aujourd_hui().isoformat()},
    )
    assert reponse.status_code == 409


def _inventaire_termine(
    api: TestClient, *, comptees: int, debut: date | None = None, commentaire: str | None = None
) -> dict[str, Any]:
    inv = _demarrer(api)
    assert api.post(f"{API}/inventaires/{inv['id']}/valider-stock").status_code == 200
    corps: dict[str, Any] = {"especes_comptees_cents": comptees, "commentaire": commentaire}
    if debut:
        corps["debut"] = debut.isoformat()
    reponse = api.post(f"{API}/inventaires/{inv['id']}/terminer", json=corps)
    assert reponse.status_code == 200, reponse.text
    return reponse.json()


def test_terminer_fige_l_ecart_especes(catalogue, admin, client) -> None:
    _vente(client, [(catalogue["the"], 3)], quand=_il_y_a(minutes=30))
    _vente(client, [(catalogue["gateau"], 1)], quand=_il_y_a(minutes=20))
    hier = buvette_crud.aujourd_hui() - timedelta(days=1)
    fini = _inventaire_termine(admin, comptees=500, debut=hier, commentaire="  Boite un peu legere  ")

    assert fini["statut"] == "termine"
    assert fini["termine_le"] is not None
    assert fini["especes_attendues_cents"] == 550
    assert fini["especes_comptees_cents"] == 500
    assert fini["ecart_especes_cents"] == -50
    assert fini["nb_ventes_especes"] == 2
    assert fini["commentaire"] == "Boite un peu legere"
    assert fini["periode_especes_debut"] is not None and fini["periode_especes_fin"] is not None
    assert admin.get(f"{API}/inventaires/en-cours").json() == {"inventaire": None}
    # Fige : l'etape ne se rejoue pas.
    assert admin.post(
        f"{API}/inventaires/{fini['id']}/terminer", json={"especes_comptees_cents": 1}
    ).status_code == 409

    # Relu apres coup : periode et attendu figes.
    especes = admin.get(f"{API}/inventaires/{fini['id']}/especes").json()
    assert especes["attendu_cents"] == 550 and especes["nb_ventes"] == 2
    assert len(especes["ventes"]) == 2


def test_la_periode_part_de_la_fin_du_dernier_inventaire_termine(
    catalogue, admin, client, db_session
) -> None:
    hier = buvette_crud.aujourd_hui() - timedelta(days=1)
    _vente(client, [(catalogue["the"], 1)], quand=_il_y_a(hours=2))
    premier = _inventaire_termine(admin, comptees=100, debut=hier)
    assert premier["especes_attendues_cents"] == 100
    # Le premier inventaire s'est termine il y a une heure.
    _deplacer(db_session, premier["id"], termine_le=_utc_naif(hours=1), periode_especes_fin=_utc_naif(hours=1))

    _vente(client, [(catalogue["gateau"], 1)], quand=_il_y_a(minutes=30))
    _vente(client, [(catalogue["chips"], 1)], carte=True, quand=_il_y_a(minutes=20))
    second = _demarrer(admin)
    # `debut` est ignore des qu'un inventaire termine existe.
    especes = admin.get(
        f"{API}/inventaires/{second['id']}/especes", params={"debut": hier.isoformat()}
    ).json()
    assert especes["premier_inventaire"] is False
    assert especes["nb_ventes"] == 1
    assert especes["attendu_cents"] == 250
    assert especes["periode_debut"] is not None

    admin.post(f"{API}/inventaires/{second['id']}/valider-stock")
    fini = admin.post(
        f"{API}/inventaires/{second['id']}/terminer", json={"especes_comptees_cents": 300}
    ).json()
    assert fini["especes_attendues_cents"] == 250
    assert fini["ecart_especes_cents"] == 50


# ---------------------------------------------------------------------------
# Droits
# ---------------------------------------------------------------------------


def test_la_compta_lit_mais_n_ecrit_pas(catalogue, admin, client_authenticated_as, compta_user) -> None:
    inv = _demarrer(admin)
    compta = client_authenticated_as(compta_user)
    assert compta.get(f"{API}/inventaires/en-cours").status_code == 200
    assert compta.get(f"{API}/inventaires/{inv['id']}").status_code == 200
    assert compta.get(f"{API}/inventaires").status_code == 200
    assert compta.get(f"{API}/inventaires/{inv['id']}/especes").status_code == 200
    assert compta.get(f"{API}/inventaires/{inv['id']}/export.xlsx").status_code == 200
    assert compta.get(f"{API}/inventaires/export.xlsx").status_code == 200
    assert compta.get(f"{API}/paiements/export.xlsx").status_code == 200

    assert compta.post(f"{API}/inventaires").status_code == 403
    assert compta.put(f"{API}/inventaires/{inv['id']}/lignes", json={"lignes": []}).status_code == 403
    assert compta.post(f"{API}/inventaires/{inv['id']}/valider-stock").status_code == 403
    assert compta.post(
        f"{API}/inventaires/{inv['id']}/terminer", json={"especes_comptees_cents": 0}
    ).status_code == 403
    assert compta.delete(f"{API}/inventaires/{inv['id']}").status_code == 403


def test_l_admin_stock_fait_l_inventaire(catalogue, client_authenticated_as, admin_stock_user) -> None:
    api = client_authenticated_as(admin_stock_user)
    fini = _inventaire_termine(api, comptees=0, debut=buvette_crud.aujourd_hui())
    assert fini["statut"] == "termine"
    assert fini["cree_par"] == admin_stock_user.full_name


def test_un_benevole_n_a_pas_acces(catalogue, client_authenticated_as, benevole_user) -> None:
    api = client_authenticated_as(benevole_user)
    assert api.get(f"{API}/inventaires").status_code == 403
    assert api.get(f"{API}/inventaires/en-cours").status_code == 403
    assert api.post(f"{API}/inventaires").status_code == 403
    assert api.get(f"{API}/paiements/export.xlsx").status_code == 403


def test_sans_jeton(client) -> None:
    assert client.get(f"{API}/inventaires").status_code == 401


# ---------------------------------------------------------------------------
# Historique
# ---------------------------------------------------------------------------


@pytest.fixture()
def historique(catalogue, admin, db_session) -> list[dict[str, Any]]:
    """Trois inventaires : 01/10 et 05/10 (termines), et un en cours aujourd'hui."""
    a = _inventaire_termine(admin, comptees=0, debut=date(2026, 9, 30))
    _deplacer(db_session, a["id"], debut_le=datetime(2026, 10, 1, 8), termine_le=datetime(2026, 10, 1, 9))
    b = _inventaire_termine(admin, comptees=0)
    _deplacer(db_session, b["id"], debut_le=datetime(2026, 10, 5, 21, 30), termine_le=datetime(2026, 10, 5, 22))
    c = _demarrer(admin)
    return [a, b, c]


def test_historique_plus_recent_d_abord(historique, admin) -> None:
    liste = admin.get(f"{API}/inventaires").json()
    assert [i["id"] for i in liste] == [historique[2]["id"], historique[1]["id"], historique[0]["id"]]
    assert "lignes" not in liste[0]
    assert set(liste[0]["resume"]) == {
        "nb_produits", "nb_ecarts", "ecart_unites", "valeur_ecart_cents", "perte_cents",
    }
    assert [i["statut"] for i in liste] == ["en_cours", "termine", "termine"]


def test_historique_filtre_par_dates_incluses(historique, admin) -> None:
    def ids(**params: str) -> list[int]:
        reponse = admin.get(f"{API}/inventaires", params=params)
        assert reponse.status_code == 200, reponse.text
        return [i["id"] for i in reponse.json()]

    a, b, _ = historique
    assert ids(debut="2026-10-01", fin="2026-10-01") == [a["id"]]
    # 21 h 30 UTC le 5 = 23 h 30 a Paris : toujours le 5.
    assert ids(debut="2026-10-02", fin="2026-10-05") == [b["id"]]
    assert ids(debut="2026-10-06", fin="2026-10-06") == []
    assert admin.get(f"{API}/inventaires", params={"debut": "2026-10-05", "fin": "2026-10-01"}).status_code == 422


# ---------------------------------------------------------------------------
# Exports Excel
# ---------------------------------------------------------------------------


def _classeur(reponse) -> Any:
    assert reponse.status_code == 200, reponse.text
    assert reponse.headers["content-type"] == XLSX
    return load_workbook(io.BytesIO(reponse.content))


def _valeurs(feuille) -> list[tuple]:
    return [tuple(c for c in ligne) for ligne in feuille.iter_rows(values_only=True)]


def test_export_d_un_inventaire(catalogue, admin, client) -> None:
    _vente(client, [(catalogue["the"], 2)], quand=_il_y_a(minutes=30))
    _vente(client, [(catalogue["gateau"], 1)], quand=_il_y_a(minutes=20))
    inv = _demarrer(admin)
    _compter(admin, inv, {"The": 8, "Gateau": 2, "Chips": 4})
    admin.post(f"{API}/inventaires/{inv['id']}/valider-stock")
    admin.post(
        f"{API}/inventaires/{inv['id']}/terminer",
        json={
            "especes_comptees_cents": 400,
            "debut": (buvette_crud.aujourd_hui() - timedelta(days=1)).isoformat(),
        },
    )

    reponse = admin.get(f"{API}/inventaires/{inv['id']}/export.xlsx")
    classeur = _classeur(reponse)
    jour = buvette_crud.aujourd_hui().isoformat()
    assert reponse.headers["content-disposition"] == f'attachment; filename="inventaire-{inv["id"]}-{jour}.xlsx"'
    assert classeur.sheetnames == ["Synthèse", "Écarts produits", "Ventes espèces"]

    synthese = {l[0]: l[1] for l in _valeurs(classeur["Synthèse"]) if l[0]}
    assert synthese["Statut"] == "Terminé"
    assert synthese["Espèces attendues"] == 4.5
    assert synthese["Espèces comptées"] == 4.0
    assert synthese["Écart espèces (compté moins attendu)"] == -0.5
    assert synthese["Perte estimée"] == 5.0
    assert isinstance(synthese["Terminé le"], datetime)

    ecarts = classeur["Écarts produits"]
    lignes = _valeurs(ecarts)
    assert lignes[0] == (
        "Produit", "Catégorie", "Prix unitaire", "Stock théorique", "Quantité comptée",
        "Écart (unités)", "Valeur de l'écart",
    )
    assert ecarts["A1"].font.bold
    assert len(lignes) == 1 + 3 + 1
    gateau = next(l for l in lignes if l[0] == "Gateau")
    assert gateau[2:] == (2.5, 4, 2, -2, -5.0)
    assert lignes[-1][0] == "Total" and lignes[-1][5] == -2 and lignes[-1][6] == -5.0
    assert ecarts.cell(row=2, column=3).number_format == "#,##0.00 €"

    ventes = _valeurs(classeur["Ventes espèces"])
    assert ventes[0] == ("Date", "Heure", "Référence", "Articles", "Montant")
    assert len(ventes) == 1 + 2 + 1
    assert ventes[1][3] == "2 × The" and ventes[1][4] == 2.0
    assert ventes[-1][0] == "Total (2 ventes)" and ventes[-1][4] == 4.5


def test_export_d_un_inventaire_en_cours(catalogue, admin) -> None:
    inv = _demarrer(admin)
    classeur = _classeur(admin.get(f"{API}/inventaires/{inv['id']}/export.xlsx"))
    assert len(_valeurs(classeur["Écarts produits"])) == 5
    assert len(_valeurs(classeur["Ventes espèces"])) == 2


def test_export_de_l_historique(historique, admin) -> None:
    reponse = admin.get(f"{API}/inventaires/export.xlsx", params={"debut": "2026-10-01", "fin": "2026-10-05"})
    classeur = _classeur(reponse)
    assert reponse.headers["content-disposition"] == 'attachment; filename="inventaires-2026-10-01_2026-10-05.xlsx"'
    assert classeur.sheetnames == ["Inventaires", "Détail"]
    resumes = _valeurs(classeur["Inventaires"])
    assert resumes[0][:4] == ("N°", "Statut", "Démarré le", "Terminé le")
    assert classeur["Inventaires"]["A1"].font.bold
    # Les inventaires du 01/10 et du 05/10, plus recent d'abord, puis les totaux.
    assert [l[0] for l in resumes[1:-1]] == [historique[1]["id"], historique[0]["id"]]
    assert resumes[1][1] == "Terminé"
    assert resumes[1][2] == datetime(2026, 10, 5, 23, 30)
    assert resumes[-1][0] == "Total (2 inventaires)"
    detail = _valeurs(classeur["Détail"])
    assert detail[0][:3] == ("N° inventaire", "Démarré le", "Produit")
    assert len(detail) == 1 + 3 * 2 + 1
    assert detail[-1][0] == "Total"

    tout = admin.get(f"{API}/inventaires/export.xlsx")
    assert tout.headers["content-disposition"] == 'attachment; filename="inventaires-tout.xlsx"'
    assert len(_valeurs(_classeur(tout)["Inventaires"])) == 1 + 3 + 1


def test_export_des_paiements(catalogue, client, client_authenticated_as, compta_user, db_session) -> None:
    jour = date(2026, 10, 5)
    _vente(client, [(catalogue["the"], 2), (catalogue["gateau"], 1)], carte=True,
           quand=datetime(2026, 10, 5, 10, 15))
    _vente(client, [(catalogue["gateau"], 2)], quand=datetime(2026, 10, 5, 16, 40))
    buvette_crud.record_sale_and_decrement(
        db_session, order_id=77, payment_id=770, item_id=1, tier_id=None, name="Jus",
        quantity_sold=1, amount_cents=200, customer={"firstName": "Yusuf", "lastName": "Martin"},
        sold_at=datetime(2026, 10, 5, 12, 5), raw_event=None,
    )
    _vente(client, [(catalogue["the"], 1)], quand=datetime(2026, 10, 4, 18, 0))
    api = client_authenticated_as(compta_user)

    reponse = api.get(f"{API}/paiements/export.xlsx", params={"debut": jour.isoformat(), "fin": jour.isoformat()})
    classeur = _classeur(reponse)
    assert reponse.headers["content-disposition"] == 'attachment; filename="paiements-2026-10-05_2026-10-05.xlsx"'
    assert classeur.sheetnames == ["Paiements"]
    feuille = classeur["Paiements"]
    lignes = _valeurs(feuille)
    assert lignes[0] == (
        "Date", "Heure", "Moyen", "Référence", "Client", "Produit", "Quantité", "Montant ligne", "Total vente",
    )
    assert feuille["A1"].font.bold
    # 2 articles carte + 1 especes + 1 HelloAsso, puis les totaux.
    assert len(lignes) == 1 + 4 + 1
    # Plus recentes d'abord, comme l'ecran.
    assert lignes[1][2] == "Espèces" and lignes[1][0] == datetime(2026, 10, 5, 16, 40)
    helloasso = next(l for l in lignes if l[2] == "HelloAsso")
    assert helloasso[3] == "Commande HelloAsso 77" and helloasso[4] == "Yusuf Martin"
    carte = [l for l in lignes if l[2] == "Carte"]
    assert len(carte) == 2 and carte[0][3].startswith("SumUp TX")
    assert carte[0][8] == 4.5
    total = lignes[-1]
    assert total[0] == "Total (3 ventes)"
    assert total[6] == 2 + 1 + 2 + 1
    assert total[7] == 4.5 + 5.0 + 2.0 and total[8] == 11.5
    assert feuille.cell(row=2, column=8).number_format == "#,##0.00 €"

    # Meme filtre de moyen que l'onglet.
    especes = _classeur(api.get(f"{API}/paiements/export.xlsx",
                                params={"debut": jour.isoformat(), "fin": jour.isoformat(), "moyen": "especes"}))
    assert len(_valeurs(especes["Paiements"])) == 1 + 1 + 1

    # Defaut : la journee en cours.
    defaut = api.get(f"{API}/paiements/export.xlsx")
    aujourd_hui = buvette_crud.aujourd_hui().isoformat()
    assert defaut.headers["content-disposition"] == (
        f'attachment; filename="paiements-{aujourd_hui}_{aujourd_hui}.xlsx"'
    )
