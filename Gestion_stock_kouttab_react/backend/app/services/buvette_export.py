"""Classeurs Excel de la buvette : inventaire, historique des inventaires,
paiements, reapprovisionnements, clotures de caisse.

Fonctions pures : elles recoivent les donnees deja calculees par le CRUD (les
memes dictionnaires que l'API sert a l'ecran) et rendent les octets du classeur.
Un export ne recalcule rien, sinon il finirait par contredire l'ecran.

Conventions : en-tetes en gras sur fond vert, montants en euros (nombres, pas
du texte, pour que le tableur puisse les additionner), horodatages a l'heure de
Paris, une ligne de totaux en bas de chaque tableau.
"""

from __future__ import annotations

import io
from datetime import date, datetime, timezone
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.crud.buvette import PARIS


MEDIA_TYPE_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

FORMAT_EUROS = "#,##0.00 €"
FORMAT_DATE = "DD/MM/YYYY"
FORMAT_DATE_HEURE = "DD/MM/YYYY HH:MM"
FORMAT_HEURE = "HH:MM"

STATUTS = {
    "en_cours": "En cours (comptage)",
    "stock_valide": "Stock validé",
    "termine": "Terminé",
}
MOYENS = {"carte": "Carte", "especes": "Espèces", "helloasso": "HelloAsso"}
ORIGINES = {"app": "Application", "tablette": "Tablette"}
TYPES_MOUVEMENT = {
    "reappro": "Réapprovisionnement",
    "vente": "Vente",
    "ecart": "Écart d'inventaire",
    "cloture": "Clôture de caisse",
}
CATEGORIES = {
    "sucre_sale": "Sucré / salé",
    "boissons": "Boissons",
    "cafe": "Café",
    "epicerie": "Épicerie",
}

_GRAS = Font(bold=True)
_ENTETE_POLICE = Font(bold=True, color="FFFFFF")
_ENTETE_FOND = PatternFill("solid", fgColor="3D4F3D")
_TOTAL_FOND = PatternFill("solid", fgColor="EFE6D6")


# ---------------------------------------------------------------------------
# Outils
# ---------------------------------------------------------------------------


def euros(cents: int | None) -> float | None:
    return None if cents is None else round(cents / 100, 2)


def heure_paris(instant: datetime | None) -> datetime | None:
    """Horodatage (UTC, avec ou sans fuseau) -> heure de Paris, sans fuseau.

    Excel ne connait pas les fuseaux : une date avec fuseau y est refusee.
    """
    if instant is None:
        return None
    if instant.tzinfo is None:
        instant = instant.replace(tzinfo=timezone.utc)
    return instant.astimezone(PARIS).replace(tzinfo=None)


def _entetes(feuille: Worksheet, ligne: int, entetes: Iterable[str]) -> None:
    for colonne, texte in enumerate(entetes, start=1):
        cellule = feuille.cell(row=ligne, column=colonne, value=texte)
        cellule.font = _ENTETE_POLICE
        cellule.fill = _ENTETE_FOND
        cellule.alignment = Alignment(wrap_text=True, vertical="center")
    feuille.freeze_panes = feuille.cell(row=ligne + 1, column=1)


def _ligne(
    feuille: Worksheet,
    ligne: int,
    valeurs: Iterable[Any],
    formats: dict[int, str] | None = None,
    *,
    total: bool = False,
) -> None:
    formats = formats or {}
    for colonne, valeur in enumerate(valeurs, start=1):
        cellule = feuille.cell(row=ligne, column=colonne, value=valeur)
        if colonne in formats and valeur is not None:
            cellule.number_format = formats[colonne]
        if total:
            cellule.font = _GRAS
            cellule.fill = _TOTAL_FOND


def _largeurs(feuille: Worksheet, largeurs: Iterable[float]) -> None:
    for colonne, largeur in enumerate(largeurs, start=1):
        feuille.column_dimensions[get_column_letter(colonne)].width = largeur


def _octets(classeur: Workbook) -> bytes:
    tampon = io.BytesIO()
    classeur.save(tampon)
    return tampon.getvalue()


def _categorie(code: str | None) -> str:
    return CATEGORIES.get(code or "", code or "")


def _articles_texte(articles: list[dict[str, Any]]) -> str:
    return ", ".join(f"{a['quantite']} × {a['nom']}" for a in articles)


# ---------------------------------------------------------------------------
# Un inventaire
# ---------------------------------------------------------------------------


ENTETES_ECARTS = (
    "Produit",
    "Catégorie",
    "Prix unitaire",
    "Stock théorique",
    "Quantité comptée",
    "Écart (unités)",
    "Valeur de l'écart",
)
ENTETES_VENTES_ESPECES = ("Date", "Heure", "Référence", "Articles", "Montant")


def _feuille_synthese(
    feuille: Worksheet, inv: dict[str, Any], mouvements: dict[str, Any] | None = None
) -> None:
    resume = inv["resume"]
    feuille.cell(row=1, column=1, value=f"Inventaire buvette n° {inv['id']}, Le Kouttâb").font = Font(
        bold=True, size=13
    )
    lignes: list[tuple[str, Any, str | None]] = [
        ("Statut", STATUTS.get(inv["statut"], inv["statut"]), None),
        ("Démarré le", heure_paris(inv["debut_le"]), FORMAT_DATE_HEURE),
        ("Stock validé le", heure_paris(inv["stock_valide_le"]), FORMAT_DATE_HEURE),
        ("Terminé le", heure_paris(inv["termine_le"]), FORMAT_DATE_HEURE),
        ("Réalisé par", inv["cree_par"] or "", None),
        ("", None, None),
        ("Produits comptés", resume["nb_produits"], None),
        ("Produits en écart", resume["nb_ecarts"], None),
        ("Écart total (unités)", resume["ecart_unites"], None),
        ("Valeur de l'écart", euros(resume["valeur_ecart_cents"]), FORMAT_EUROS),
        ("Perte estimée", euros(resume["perte_cents"]), FORMAT_EUROS),
        ("", None, None),
        *_lignes_achats(mouvements),
        *_lignes_carte(mouvements),
        *_lignes_clotures(mouvements),
        ("Période des espèces : du", heure_paris(inv["periode_especes_debut"]), FORMAT_DATE_HEURE),
        ("Période des espèces : au", heure_paris(inv["periode_especes_fin"]), FORMAT_DATE_HEURE),
        ("Ventes en espèces", inv["nb_ventes_especes"], None),
        ("Espèces attendues", euros(inv["especes_attendues_cents"]), FORMAT_EUROS),
        ("Espèces comptées", euros(inv["especes_comptees_cents"]), FORMAT_EUROS),
        ("Écart espèces (compté moins attendu)", euros(inv["ecart_especes_cents"]), FORMAT_EUROS),
        ("Commentaire", inv["commentaire"] or "", None),
    ]
    for rang, (libelle, valeur, fmt) in enumerate(lignes, start=3):
        if not libelle:
            continue
        feuille.cell(row=rang, column=1, value=libelle).font = _GRAS
        cellule = feuille.cell(row=rang, column=2, value=valeur)
        if fmt and valeur is not None:
            cellule.number_format = fmt
        cellule.alignment = Alignment(horizontal="left", wrap_text=True)
    _largeurs(feuille, (38, 44))


def _lignes_achats(mouvements: dict[str, Any] | None) -> list[tuple[str, Any, str | None]]:
    """Lignes de la synthese consacrees aux mouvements et aux achats de la periode."""
    if mouvements is None:
        return []
    reappros = mouvements["reappros"]
    debut = heure_paris(mouvements["periode_debut"])
    return [
        (
            "Période des mouvements : du",
            debut if debut is not None else "Début des données",
            FORMAT_DATE_HEURE if debut is not None else None,
        ),
        ("Période des mouvements : au", heure_paris(mouvements["periode_fin"]), FORMAT_DATE_HEURE),
        ("Réapprovisionnements", len(reappros), None),
        ("Unités réapprovisionnées", sum(r["quantite"] for r in reappros), None),
        ("Total des achats", euros(mouvements["achats_cents"]), FORMAT_EUROS),
        ("", None, None),
    ]


def _lignes_carte(mouvements: dict[str, Any] | None) -> list[tuple[str, Any, str | None]]:
    """Ventes carte de la periode des mouvements : brut, frais SumUp, net."""
    if mouvements is None or mouvements.get("ventes_carte") is None:
        return []
    carte = mouvements["ventes_carte"]
    return [
        ("Ventes carte de la période", carte["nb"], None),
        ("Ventes carte (brut)", euros(carte["brut_cents"]), FORMAT_EUROS),
        ("Frais SumUp", euros(carte["frais_cents"]), FORMAT_EUROS),
        ("Ventes carte (net)", euros(carte["net_cents"]), FORMAT_EUROS),
        ("", None, None),
    ]


def _lignes_clotures(mouvements: dict[str, Any] | None) -> list[tuple[str, Any, str | None]]:
    """Lignes de la synthese consacrees aux clotures de caisse de la periode."""
    if mouvements is None:
        return []
    clotures = mouvements.get("clotures") or []
    return [
        ("Clôtures de caisse de la période", len(clotures), None),
        (
            "Cumul des écarts des clôtures",
            euros(sum(c["ecart_cents"] for c in clotures)),
            FORMAT_EUROS,
        ),
        ("", None, None),
    ]


def _feuille_ecarts(feuille: Worksheet, lignes: list[dict[str, Any]]) -> None:
    _entetes(feuille, 1, ENTETES_ECARTS)
    formats = {3: FORMAT_EUROS, 7: FORMAT_EUROS}
    rang = 2
    for l in lignes:
        _ligne(
            feuille,
            rang,
            (
                l["nom"],
                _categorie(l["categorie"]),
                euros(l["prix_cents"]),
                l["quantite_theorique"],
                l["quantite_comptee"],
                l["ecart"],
                euros(l["valeur_ecart_cents"]),
            ),
            formats,
        )
        rang += 1
    _ligne(
        feuille,
        rang,
        (
            "Total",
            "",
            None,
            sum(l["quantite_theorique"] or 0 for l in lignes),
            sum(l["quantite_comptee"] for l in lignes),
            sum(l["ecart"] or 0 for l in lignes),
            euros(sum(l["valeur_ecart_cents"] or 0 for l in lignes)),
        ),
        formats,
        total=True,
    )
    _largeurs(feuille, (32, 16, 13, 15, 16, 14, 17))


def _feuille_ventes_especes(feuille: Worksheet, ventes: list[dict[str, Any]]) -> None:
    _entetes(feuille, 1, ENTETES_VENTES_ESPECES)
    rang = 2
    for v in ventes:
        instant = heure_paris(v["sold_at"])
        _ligne(
            feuille,
            rang,
            (instant, instant, v["cle"], _articles_texte(v["articles"]), euros(v["total_cents"])),
            {1: FORMAT_DATE, 2: FORMAT_HEURE, 5: FORMAT_EUROS},
        )
        rang += 1
    _ligne(
        feuille,
        rang,
        (f"Total ({len(ventes)} ventes)", None, "", "", euros(sum(v["total_cents"] for v in ventes))),
        {5: FORMAT_EUROS},
        total=True,
    )
    _largeurs(feuille, (14, 8, 40, 60, 12))


ENTETES_MOUVEMENTS = (
    "Date",
    "Heure",
    "Type",
    "Produit",
    "Quantité",
    "Moyen ou origine",
    "Prix unitaire",
    "Montant",
    "Frais SumUp",
    "Détail",
)
ENTETES_RECAP = (
    "Produit",
    "Catégorie",
    "Prix unitaire",
    "Compté au précédent inventaire",
    "+ Réappros",
    "− Ventes",
    "= Stock attendu",
    "Compté",
    "Écart (unités)",
    "Valeur de l'écart",
)


def _feuille_mouvements(feuille: Worksheet, mouvements: list[dict[str, Any]]) -> None:
    """Chronologie unique : reappros, ventes par article, clotures de caisse
    (montant compte), ecarts d'inventaire."""
    _entetes(feuille, 1, ENTETES_MOUVEMENTS)
    formats = {1: FORMAT_DATE, 2: FORMAT_HEURE, 7: FORMAT_EUROS, 8: FORMAT_EUROS, 9: FORMAT_EUROS}
    rang = 2
    for m in mouvements:
        if m["type"] == "vente":
            moyen = MOYENS.get(m["moyen"], m["moyen"])
        elif m["type"] == "reappro":
            moyen = ORIGINES.get(m["moyen"], m["moyen"])
        elif m["type"] == "cloture":
            moyen = MOYENS["especes"]
        else:
            moyen = ""
        _ligne(
            feuille,
            rang,
            (
                m["quand"],
                m["quand"],
                TYPES_MOUVEMENT.get(m["type"], m["type"]),
                m["produit"],
                m["quantite"],
                moyen,
                euros(m["prix_unitaire_cents"]),
                euros(m["montant_cents"]),
                euros(m.get("frais_cents")),
                m["detail"] or "",
            ),
            formats,
        )
        rang += 1
    _ligne(
        feuille,
        rang,
        (
            f"Total ({len(mouvements)} mouvements)",
            None,
            "",
            "",
            sum(m["quantite"] or 0 for m in mouvements),
            "",
            None,
            None,
            euros(sum(m.get("frais_cents") or 0 for m in mouvements)),
            "",
        ),
        formats,
        total=True,
    )
    _largeurs(feuille, (12, 8, 22, 30, 10, 16, 13, 13, 13, 40))


def _feuille_recap(feuille: Worksheet, recap: list[dict[str, Any]]) -> None:
    """Par produit : compte au precedent + reappros - ventes = attendu, compare au compte."""
    _entetes(feuille, 1, ENTETES_RECAP)
    formats = {3: FORMAT_EUROS, 10: FORMAT_EUROS}
    rang = 2
    for r in recap:
        _ligne(
            feuille,
            rang,
            (
                r["nom"],
                _categorie(r["categorie"]),
                euros(r["prix_cents"]),
                r["compte_precedent"],
                r["reappros"],
                r["ventes"],
                r["attendu"],
                r["compte"],
                r["ecart"],
                euros(r["valeur_ecart_cents"]),
            ),
            formats,
        )
        rang += 1
    _ligne(
        feuille,
        rang,
        (
            "Total",
            "",
            None,
            sum(r["compte_precedent"] or 0 for r in recap),
            sum(r["reappros"] for r in recap),
            sum(r["ventes"] for r in recap),
            sum(r["attendu"] or 0 for r in recap),
            sum(r["compte"] for r in recap),
            sum(r["ecart"] or 0 for r in recap),
            euros(sum(r["valeur_ecart_cents"] or 0 for r in recap)),
        ),
        formats,
        total=True,
    )
    feuille.cell(
        row=rang + 2,
        column=1,
        value=(
            "Case vide : quantité inconnue au précédent inventaire (produit non compté, "
            "ou premier inventaire). L'attendu et l'écart ne sont alors pas calculés."
        ),
    ).font = Font(italic=True)
    _largeurs(feuille, (32, 16, 13, 16, 12, 12, 14, 11, 13, 16))


def classeur_inventaire(
    inv: dict[str, Any], especes: dict[str, Any], mouvements: dict[str, Any] | None = None
) -> bytes:
    """Feuilles « Synthèse », « Écarts produits », « Ventes espèces » ; avec
    `mouvements` (`crud.buvette_reappro.mouvements_inventaire`), en plus
    « Réapprovisionnements », « Mouvements », « Récap par produit » et
    « Clôtures de caisse » (celles de la période des mouvements)."""
    classeur = Workbook()
    synthese = classeur.active
    synthese.title = "Synthèse"
    _feuille_synthese(synthese, inv, mouvements)
    _feuille_ecarts(classeur.create_sheet("Écarts produits"), inv["lignes"])
    _feuille_ventes_especes(classeur.create_sheet("Ventes espèces"), especes["ventes"])
    if mouvements is not None:
        _feuille_reappros(classeur.create_sheet("Réapprovisionnements"), mouvements["reappros"])
        _feuille_mouvements(classeur.create_sheet("Mouvements"), mouvements["mouvements"])
        _feuille_recap(classeur.create_sheet("Récap par produit"), mouvements["recap"])
        _feuille_clotures(
            classeur.create_sheet("Clôtures de caisse"), mouvements.get("clotures") or []
        )
    return _octets(classeur)


# ---------------------------------------------------------------------------
# Historique des inventaires
# ---------------------------------------------------------------------------


ENTETES_HISTORIQUE = (
    "N°",
    "Statut",
    "Démarré le",
    "Terminé le",
    "Réalisé par",
    "Produits comptés",
    "Produits en écart",
    "Écart (unités)",
    "Valeur de l'écart",
    "Perte estimée",
    "Ventes en espèces",
    "Espèces attendues",
    "Espèces comptées",
    "Écart espèces",
    "Commentaire",
)
ENTETES_DETAIL = (
    "N° inventaire",
    "Démarré le",
    "Produit",
    "Catégorie",
    "Prix unitaire",
    "Stock théorique",
    "Quantité comptée",
    "Écart (unités)",
    "Valeur de l'écart",
)


def classeur_historique(
    inventaires: list[dict[str, Any]],
    reappros: list[dict[str, Any]] | None = None,
    clotures: list[dict[str, Any]] | None = None,
) -> bytes:
    """Feuilles « Inventaires » (un résumé par ligne), « Détail » (les lignes
    produits) et, si fournis, « Réapprovisionnements » et « Clôtures de caisse »
    de la période filtrée."""
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Inventaires"
    _entetes(feuille, 1, ENTETES_HISTORIQUE)
    formats = {3: FORMAT_DATE_HEURE, 4: FORMAT_DATE_HEURE, 9: FORMAT_EUROS, 10: FORMAT_EUROS,
               12: FORMAT_EUROS, 13: FORMAT_EUROS, 14: FORMAT_EUROS}
    rang = 2
    for inv in inventaires:
        r = inv["resume"]
        _ligne(
            feuille,
            rang,
            (
                inv["id"],
                STATUTS.get(inv["statut"], inv["statut"]),
                heure_paris(inv["debut_le"]),
                heure_paris(inv["termine_le"]),
                inv["cree_par"] or "",
                r["nb_produits"],
                r["nb_ecarts"],
                r["ecart_unites"],
                euros(r["valeur_ecart_cents"]),
                euros(r["perte_cents"]),
                inv["nb_ventes_especes"],
                euros(inv["especes_attendues_cents"]),
                euros(inv["especes_comptees_cents"]),
                euros(inv["ecart_especes_cents"]),
                inv["commentaire"] or "",
            ),
            formats,
        )
        rang += 1
    _ligne(
        feuille,
        rang,
        (
            f"Total ({len(inventaires)} inventaires)",
            "",
            None,
            None,
            "",
            sum(i["resume"]["nb_produits"] for i in inventaires),
            sum(i["resume"]["nb_ecarts"] for i in inventaires),
            sum(i["resume"]["ecart_unites"] for i in inventaires),
            euros(sum(i["resume"]["valeur_ecart_cents"] for i in inventaires)),
            euros(sum(i["resume"]["perte_cents"] for i in inventaires)),
            sum(i["nb_ventes_especes"] or 0 for i in inventaires),
            euros(sum(i["especes_attendues_cents"] or 0 for i in inventaires)),
            euros(sum(i["especes_comptees_cents"] or 0 for i in inventaires)),
            euros(sum(i["ecart_especes_cents"] or 0 for i in inventaires)),
            "",
        ),
        formats,
        total=True,
    )
    _largeurs(feuille, (22, 18, 17, 17, 22, 11, 11, 11, 14, 14, 11, 14, 14, 14, 40))

    detail = classeur.create_sheet("Détail")
    _entetes(detail, 1, ENTETES_DETAIL)
    formats_detail = {2: FORMAT_DATE_HEURE, 5: FORMAT_EUROS, 9: FORMAT_EUROS}
    rang = 2
    toutes = [(inv, l) for inv in inventaires for l in inv["lignes"]]
    for inv, l in toutes:
        _ligne(
            detail,
            rang,
            (
                inv["id"],
                heure_paris(inv["debut_le"]),
                l["nom"],
                _categorie(l["categorie"]),
                euros(l["prix_cents"]),
                l["quantite_theorique"],
                l["quantite_comptee"],
                l["ecart"],
                euros(l["valeur_ecart_cents"]),
            ),
            formats_detail,
        )
        rang += 1
    _ligne(
        detail,
        rang,
        (
            "Total",
            None,
            "",
            "",
            None,
            sum(l["quantite_theorique"] or 0 for _, l in toutes),
            sum(l["quantite_comptee"] for _, l in toutes),
            sum(l["ecart"] or 0 for _, l in toutes),
            euros(sum(l["valeur_ecart_cents"] or 0 for _, l in toutes)),
        ),
        formats_detail,
        total=True,
    )
    _largeurs(detail, (14, 17, 32, 16, 13, 15, 16, 14, 17))
    if reappros is not None:
        _feuille_reappros(classeur.create_sheet("Réapprovisionnements"), reappros)
    if clotures is not None:
        _feuille_clotures(classeur.create_sheet("Clôtures de caisse"), clotures)
    return _octets(classeur)


# ---------------------------------------------------------------------------
# Paiements
# ---------------------------------------------------------------------------


ENTETES_PAIEMENTS = (
    "Date",
    "Heure",
    "Moyen",
    "Référence",
    "Client",
    "Produit",
    "Quantité",
    "Montant ligne",
    "Total vente",
    "Frais SumUp",
    "Net vente",
)


def _reference(vente: dict[str, Any]) -> str:
    if vente.get("sumup_tx_code"):
        return f"SumUp {vente['sumup_tx_code']}"
    if vente.get("helloasso_order_id") is not None:
        return f"Commande HelloAsso {vente['helloasso_order_id']}"
    return vente["cle"]


def classeur_paiements(
    ventes: list[dict[str, Any]], totaux: dict[str, Any] | None = None
) -> bytes:
    """Feuille « Paiements » : une ligne par article, puis une ligne de totaux
    et, avec `totaux`, le detail brut / frais SumUp / net.

    `ventes` et `totaux` : ceux de `crud.buvette.paiements` (heure murale de
    Paris, sans fuseau), dans l'ordre de l'ecran. Frais et net sont ceux de la
    vente (calcules par transaction), repetes comme le total de la vente ;
    vides hors carte pour les frais.
    """
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Paiements"
    _entetes(feuille, 1, ENTETES_PAIEMENTS)
    formats = {
        1: FORMAT_DATE,
        2: FORMAT_HEURE,
        8: FORMAT_EUROS,
        9: FORMAT_EUROS,
        10: FORMAT_EUROS,
        11: FORMAT_EUROS,
    }
    rang = 2
    for v in ventes:
        instant = v["sold_at"]
        if isinstance(instant, datetime) and instant.tzinfo is not None:
            instant = heure_paris(instant)
        for a in v["articles"]:
            _ligne(
                feuille,
                rang,
                (
                    instant,
                    instant,
                    MOYENS.get(v["moyen"], v["moyen"]),
                    _reference(v),
                    v.get("client") or "",
                    a["nom"],
                    a["quantite"],
                    euros(a["montant_cents"]),
                    euros(v["total_cents"]),
                    euros(v.get("frais_cents")),
                    euros(v.get("net_cents", v["total_cents"])),
                ),
                formats,
            )
            rang += 1
    _ligne(
        feuille,
        rang,
        (
            f"Total ({len(ventes)} ventes)",
            None,
            "",
            "",
            "",
            "",
            sum(a["quantite"] for v in ventes for a in v["articles"]),
            euros(sum(a["montant_cents"] for v in ventes for a in v["articles"])),
            euros(sum(v["total_cents"] for v in ventes)),
            euros(sum(v.get("frais_cents") or 0 for v in ventes)),
            euros(sum(v.get("net_cents", v["total_cents"]) for v in ventes)),
        ),
        formats,
        total=True,
    )
    if totaux is not None:
        rang += 2
        for libelle, cle in (
            ("Carte (brut)", "carte_cents"),
            ("Frais SumUp", "frais_carte_cents"),
            ("Carte (net)", "carte_net_cents"),
            ("Espèces", "especes_cents"),
            ("HelloAsso", "helloasso_cents"),
            ("Total brut", "total_cents"),
            ("Total net encaissé", "net_total_cents"),
        ):
            feuille.cell(row=rang, column=1, value=libelle).font = _GRAS
            cellule = feuille.cell(row=rang, column=2, value=euros(totaux.get(cle, 0)))
            cellule.number_format = FORMAT_EUROS
            rang += 1
    _largeurs(feuille, (12, 8, 11, 30, 22, 30, 10, 14, 13, 13, 13))
    return _octets(classeur)


# ---------------------------------------------------------------------------
# Reapprovisionnements
# ---------------------------------------------------------------------------


ENTETES_REAPPROS = (
    "Date",
    "Heure",
    "Produit",
    "Quantité",
    "Prix unitaire",
    "Total",
    "Origine",
    "Par",
    "Stock avant",
    "Stock après",
    "Commentaire",
)


def _feuille_reappros(feuille: Worksheet, reappros: list[dict[str, Any]]) -> None:
    """Une ligne par reappro (dans l'ordre recu), puis une ligne de totaux.

    `reappros` : dictionnaires de `crud.buvette_reappro.reappro_out`.
    """
    _entetes(feuille, 1, ENTETES_REAPPROS)
    formats = {1: FORMAT_DATE, 2: FORMAT_HEURE, 5: FORMAT_EUROS, 6: FORMAT_EUROS}
    rang = 2
    for r in reappros:
        instant = heure_paris(r["created_at"])
        _ligne(
            feuille,
            rang,
            (
                instant,
                instant,
                r["nom"],
                r["quantite"],
                euros(r["prix_achat_unitaire_cents"]),
                euros(r["total_cents"]),
                ORIGINES.get(r["origine"], r["origine"]),
                r["fait_par"] or "",
                r["stock_avant"],
                r["stock_apres"],
                r["commentaire"] or "",
            ),
            formats,
        )
        rang += 1
    _ligne(
        feuille,
        rang,
        (
            f"Total ({len(reappros)} réapprovisionnements)",
            None,
            "",
            sum(r["quantite"] for r in reappros),
            None,
            euros(sum(r["total_cents"] or 0 for r in reappros)),
            "",
            "",
            None,
            None,
            "",
        ),
        formats,
        total=True,
    )
    _largeurs(feuille, (12, 8, 30, 10, 13, 13, 12, 22, 11, 11, 40))


def classeur_reappros(reappros: list[dict[str, Any]]) -> bytes:
    """Feuille « Réapprovisionnements » : l'historique filtré de l'écran."""
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Réapprovisionnements"
    _feuille_reappros(feuille, reappros)
    return _octets(classeur)


# ---------------------------------------------------------------------------
# Clotures de caisse
# ---------------------------------------------------------------------------


ENTETES_CLOTURES = (
    "Date",
    "Heure",
    "Période : du",
    "Période : au",
    "Ventes en espèces",
    "Attendu",
    "Compté",
    "Écart (compté moins attendu)",
    "Commentaire",
    "Par",
)


def _feuille_clotures(feuille: Worksheet, clotures: list[dict[str, Any]]) -> None:
    """Une ligne par cloture (dans l'ordre recu), puis une ligne de totaux.

    `clotures` : dictionnaires de `crud.buvette_cloture.cloture_out`.
    """
    _entetes(feuille, 1, ENTETES_CLOTURES)
    formats = {
        1: FORMAT_DATE,
        2: FORMAT_HEURE,
        3: FORMAT_DATE_HEURE,
        4: FORMAT_DATE_HEURE,
        6: FORMAT_EUROS,
        7: FORMAT_EUROS,
        8: FORMAT_EUROS,
    }
    rang = 2
    for c in clotures:
        instant = heure_paris(c["periode_fin"])
        _ligne(
            feuille,
            rang,
            (
                instant,
                instant,
                heure_paris(c["periode_debut"]),
                instant,
                c["nb_ventes"],
                euros(c["attendu_cents"]),
                euros(c["compte_cents"]),
                euros(c["ecart_cents"]),
                c["commentaire"] or "",
                c["saisi_par"] or "",
            ),
            formats,
        )
        rang += 1
    _ligne(
        feuille,
        rang,
        (
            f"Total ({len(clotures)} clôtures)",
            None,
            None,
            None,
            sum(c["nb_ventes"] for c in clotures),
            euros(sum(c["attendu_cents"] for c in clotures)),
            euros(sum(c["compte_cents"] for c in clotures)),
            euros(sum(c["ecart_cents"] for c in clotures)),
            "",
            "",
        ),
        formats,
        total=True,
    )
    _largeurs(feuille, (20, 8, 17, 17, 11, 12, 12, 16, 40, 22))


def classeur_clotures(clotures: list[dict[str, Any]]) -> bytes:
    """Feuille « Clôtures de caisse » : l'historique filtré de l'écran."""
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Clôtures de caisse"
    _feuille_clotures(feuille, clotures)
    return _octets(classeur)


def nom_periode(debut: date | None, fin: date | None) -> str:
    """`2026-10-01_2026-10-31`, ou la borne connue, ou `tout`."""
    if debut and fin:
        return f"{debut.isoformat()}_{fin.isoformat()}"
    if debut:
        return f"depuis-{debut.isoformat()}"
    if fin:
        return f"jusqu-au-{fin.isoformat()}"
    return "tout"
